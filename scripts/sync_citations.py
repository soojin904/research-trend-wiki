"""
research_metadata.xlsx → citations.xlsx 동기화
citations.xlsx 컬럼: title, authors, journal, publisher, year, doi,
  Web of Science Core Collection, 초록, openalex_id, kci_id,
  cited_by_count, cite_kci, cite_wos, fwci, 인용합, 구분
"""
import sys
from pathlib import Path

import openpyxl

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

WIKI_DIR = Path(__file__).parent.parent
META_PATH = WIKI_DIR / "nm-reference" / "research_metadata.xlsx"
CITE_PATH = WIKI_DIR / "nm-reference" / "citations.xlsx"

# research_metadata 컬럼 인덱스 (0-based)
M = {
    "title":    1,   # col2
    "authors":  3,   # col4
    "journal":  5,   # col6
    "publisher":6,   # col7
    "year":     7,   # col8
    "doi":      18,  # col19
    "wos":      36,  # col37
    "abstract": 20,  # col21
    "openalex_id": 22,  # col23
    "kci_id":   21,  # col22 (article_id)
    "cited_by_count": 23,  # col24
    "cite_kci": 24,  # col25
    "cite_wos": 25,  # col26
    "fwci":     26,  # col27
    "인용합":   30,  # col31
    "구분":     29,  # col30
}

CITE_HEADERS = [
    "title", "authors", "journal", "publisher", "year", "doi",
    "Web of Science Core Collection",
    "초록", "openalex_id", "kci_id",
    "cited_by_count", "cite_kci", "cite_wos", "fwci", "인용합", "구분",
]

print("=== research_metadata.xlsx 읽기 ===")
wb_m = openpyxl.load_workbook(META_PATH, read_only=True, data_only=True)
ws_m = wb_m.worksheets[0]

rows_out = []
for row in ws_m.iter_rows(min_row=2, values_only=True):
    title = row[M["title"]]
    if not title:
        continue
    rows_out.append([
        title,
        row[M["authors"]],
        row[M["journal"]],
        row[M["publisher"]],
        row[M["year"]],
        row[M["doi"]],
        row[M["wos"]],
        row[M["abstract"]],
        row[M["openalex_id"]],
        row[M["kci_id"]],
        row[M["cited_by_count"]],
        row[M["cite_kci"]],
        row[M["cite_wos"]],
        row[M["fwci"]],
        row[M["인용합"]],
        row[M["구분"]],
    ])

wb_m.close()
print(f"  추출: {len(rows_out)}건")

# year(index 4) 기준 내림차순 정렬
rows_out.sort(key=lambda r: int(r[4]) if r[4] and str(r[4]).isdigit() else 0, reverse=True)

print("\n=== citations.xlsx 쓰기 ===")
wb_c = openpyxl.Workbook()
ws_c = wb_c.active
ws_c.title = "citations"
ws_c.append(CITE_HEADERS)
for r in rows_out:
    ws_c.append(r)

wb_c.save(CITE_PATH)
print(f"  저장 완료: {CITE_PATH}")
print(f"  총 {len(rows_out)}행 / {len(CITE_HEADERS)}열")
